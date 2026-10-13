"""Ekspor Excel: laporan lengkap (rumus hidup) dan unduhan data terfilter (SUBTOTAL).

Setiap rumus juga menyimpan nilai hasilnya (disisipkan ke XML setelah disimpan), sehingga angka tampil benar
di penampil yang tidak menghitung ulang. Excel/LibreOffice tetap menghitung ulang saat file dibuka.
"""
import re
import zipfile
from datetime import date
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import inti

BARIS_MAKS = 3000
KUNING = PatternFill("solid", fgColor="FFF2CC")
ABU = PatternFill("solid", fgColor="EDEDED")
TIPIS = Side(style="thin", color="BFBFBF")
GARIS = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
F_ANGKA = "#,##0"
F_PERSEN = "0.00%"
F_CAPAIAN = "0%"
WARNA_STATUS = {"Tercapai": "C6EFCE", "Sesuai laju": "DDEBF7", "Hampir": "FFEB9C", "Perlu dikejar": "FCE4D6",
                "Tertinggal": "FFC7CE", "Belum mulai": "EDEDED"}
LEMBAR_PL = {"ig": "Instagram", "tt": "TikTok", "fb": "Facebook", "yt": "YouTube", "web": "Berita"}


class Buku:
    """Workbook openpyxl + catatan nilai hasil rumus untuk disisipkan saat disimpan."""

    def __init__(self, warna="C00000"):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.cache = {}
        self.warna = (warna or "C00000").lstrip("#").upper()

    def lembar(self, nama):
        return self.wb.create_sheet(nama[:31])

    def rumus(self, ws, sel, f, nilai, fmt=None):
        ws[sel] = f
        if fmt:
            ws[sel].number_format = fmt
        self.cache[(ws.title, sel)] = nilai
        return ws[sel]

    def kepala(self, ws, baris, judul, lebar=None):
        for i, j in enumerate(judul, start=1):
            c = ws.cell(row=baris, column=i, value=j)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=self.warna)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = GARIS
        if lebar:
            for i, w in enumerate(lebar, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w

    def warnai_status(self, ws, rentang):
        for teks, warna in WARNA_STATUS.items():
            ws.conditional_formatting.add(rentang, CellIsRule(operator="equal", formula=[f'"{teks}"'],
                                                              fill=PatternFill("solid", bgColor=warna)))

    def simpan(self):
        buf = BytesIO()
        self.wb.save(buf)
        return sisipkan_nilai(buf.getvalue(), [ws.title for ws in self.wb.worksheets], self.cache)


def _v(nilai):
    if nilai is None or nilai == "":
        return None, None
    if isinstance(nilai, bool):
        return "b", "1" if nilai else "0"
    if isinstance(nilai, (int, float)):
        return None, repr(float(nilai)) if isinstance(nilai, float) else str(nilai)
    return "str", escape(str(nilai))


def sisipkan_nilai(xlsx, judul_lembar, cache):
    per_lembar = {}
    for (ws, sel), v in cache.items():
        per_lembar.setdefault(ws, {})[sel] = v
    zin = zipfile.ZipFile(BytesIO(xlsx))
    out = BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    pola = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)
    for item in zin.infolist():
        data = zin.read(item.filename)
        m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
        if m and int(m.group(1)) - 1 < len(judul_lembar):
            nilai = per_lembar.get(judul_lembar[int(m.group(1)) - 1], {})
            if nilai:
                def ganti(mm):
                    t, v = _v(nilai.get(mm.group(1)))
                    attr = re.sub(r'\st="[^"]*"', "", mm.group(2))
                    if v is None:
                        return f'<c r="{mm.group(1)}"{attr}><f>{mm.group(3)}</f></c>'
                    tt = f' t="{t}"' if t else ""
                    return f'<c r="{mm.group(1)}"{attr}{tt}><f>{mm.group(3)}</f><v>{v}</v></c>'
                data = pola.sub(ganti, data.decode("utf-8")).encode("utf-8")
        zout.writestr(item, data)
    zout.close()
    return out.getvalue()


def _status_rumus(cap, porsi_ref, berjalan, minimum):
    if not berjalan or minimum:
        return f'IF({cap}="","",IF({cap}>=1,"Tercapai",IF({cap}>={inti.AMBANG_HAMPIR},"Hampir","Tertinggal")))'
    return (f'IF({cap}="","",IF({cap}>=1,"Tercapai",IF({porsi_ref}=0,"Belum mulai",'
            f'IF({cap}/{porsi_ref}>={inti.AMBANG_LAJU},"Sesuai laju",IF({cap}/{porsi_ref}>={inti.AMBANG_KEJAR},"Perlu dikejar","Tertinggal")))))')


def _lbl(st):
    return inti.STATUS_LABEL.get(st, "") if st else ""


def _tampil_posts(posts):
    """Salinan postingan dengan nama kelompok/akun memakai penulisan baku (sama dengan pengelompokan di layar)."""
    tk, ta = inti.nama_tampil(posts, "kelompok"), inti.nama_tampil(posts, "akun")
    out = []
    for p in posts:
        q = dict(p)
        q["kelompok"] = tk.get(inti.name_key(p.get("kelompok")), "")
        q["akun"] = ta.get(inti.name_key(p.get("akun")), "")
        out.append(q)
    return out


def _urut(posts):
    return sorted(posts, key=lambda p: (p.get("tanggal") or "9999", (p.get("kelompok") or "").lower(), (p.get("akun") or "").lower()))


def tulis_lembar_platform(bk, pl, rows, label_kelompok):
    ws = bk.lembar(LEMBAR_PL[pl])
    judul = ["No", label_kelompok, "Tanggal", "Nama Akun", "Jenis", "Judul", "Link", "Views", "Like", "Comment",
             "Repost", "Share", "Total", "ER"]
    bk.kepala(ws, 1, judul, [5, 18, 12, 20, 9, 50, 40, 11, 9, 9, 9, 9, 10, 9])
    for i, p in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=p.get("kelompok") or None)
        if p.get("tanggal"):
            c = ws.cell(row=i, column=3, value=date.fromisoformat(p["tanggal"]))
            c.number_format = "dd/mm/yyyy"
        ws.cell(row=i, column=4, value=p.get("akun") or None)
        ws.cell(row=i, column=5, value=p.get("jenis"))
        ws.cell(row=i, column=6, value=p.get("judul") or None)
        c = ws.cell(row=i, column=7, value=p.get("link") or "(link kosong)")
        if p.get("link"):
            info = inti.parse_link(p["link"])
            if info:
                c.hyperlink = info["url"]
                c.font = Font(color="0563C1", underline="single")
        if pl != "web":
            for j, m in enumerate(inti.METRIK, start=8):
                c = ws.cell(row=i, column=j, value=p.get(m))
                c.number_format = F_ANGKA
            bk.rumus(ws, f"M{i}", f"=SUM(I{i}:L{i})", inti.interaksi(p), F_ANGKA)
            er = inti.er_baris(p)
            bk.rumus(ws, f"N{i}", f'=IFERROR(M{i}/H{i},"")', er if er is not None else "", F_PERSEN)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:N{max(2, len(rows) + 1)}"
    return ws


def buat_laporan(klien, r, posts, flags):
    """r = inti.ringkas(...) untuk posts yang sama."""
    bk = Buku(klien.get("warna"))
    label = klien.get("label_kelompok") or "Kelompok"
    pakai_kel = klien.get("pakai_kelompok", True)
    posts_t = _tampil_posts(posts)
    berjalan = r["berjalan"]
    porsi = r["porsi"] if r["porsi"] is not None else 1
    rng = {pl: f"'{LEMBAR_PL[pl]}'!" for pl in inti.PLATFORM}
    R = BARIS_MAKS

    # ---------------- Ringkasan
    ws = bk.lembar("Ringkasan")
    ws["A1"] = f"Laporan Publikasi — {klien['nama']}"
    ws["A1"].font = Font(bold=True, size=16, color=bk.warna)
    ws["A2"], ws["B2"] = "Klien", klien["nama"]
    ws["A3"], ws["B3"] = "Periode", r["nama_periode"]
    ws["A4"], ws["B4"] = "Porsi hari berlalu", porsi
    ws["B4"].number_format = F_CAPAIAN
    ws["C4"] = "periode berjalan: status memakai laju harian" if berjalan else "periode sudah lewat"
    ws["C4"].font = Font(italic=True, color="7F7F7F")
    bk.kepala(ws, 6, ["Indikator", "Realisasi", "Target", "Capaian", "Status", "Bulan lalu", "Perubahan"],
              [30, 14, 14, 11, 15, 14, 12])
    baris_ind = {}
    for n, ind in enumerate(r["indikator"], start=7):
        baris_ind[ind["kunci"]] = n
    for ind in r["indikator"]:
        n = baris_ind[ind["kunci"]]
        k = ind["kunci"]
        ws.cell(row=n, column=1, value=ind["label"])
        is_er = ind["jenis"] == "er"
        fmt = F_PERSEN if is_er else F_ANGKA
        if k.startswith("posts."):
            pl = k[6:]
            f = f"=COUNTA({rng[pl]}G2:G{R})"
        elif k == "publikasi":
            f = f"=SUM(B{baris_ind['posts.ig']}:B{baris_ind['posts.web']})"
        elif k == "views":
            f = "=SUM(" + ",".join(f"{rng[pl]}H2:H{R}" for pl in inti.SOSMED) + ")"
        elif k == "interaksi":
            f = "=SUM(" + ",".join(f"{rng[pl]}I2:L{R}" for pl in inti.SOSMED) + ")"
        elif k == "er":
            f = f'=IFERROR(B{baris_ind["interaksi"]}/B{baris_ind["views"]},"")'
        else:
            pl = k[3:]
            f = f'=IFERROR(SUM({rng[pl]}I2:L{R})/SUM({rng[pl]}H2:H{R}),"")'
        bk.rumus(ws, f"B{n}", f, ind["realisasi"] if ind["realisasi"] is not None else "", fmt)
        if k == "publikasi":
            bk.rumus(ws, f"C{n}", f'=IF(COUNT(C{baris_ind["posts.ig"]}:C{baris_ind["posts.web"]})=0,"",'
                                  f'SUM(C{baris_ind["posts.ig"]}:C{baris_ind["posts.web"]}))', ind["target"] or "", fmt)
        else:
            c = ws.cell(row=n, column=3, value=ind["target"])
            c.number_format = fmt
            c.fill = KUNING
        bk.rumus(ws, f"D{n}", f'=IF(OR(N(C{n})=0,B{n}=""),"",B{n}/C{n})',
                 ind["capaian"] if ind["capaian"] is not None else "", F_CAPAIAN)
        bk.rumus(ws, f"E{n}", "=" + _status_rumus(f"D{n}", "$B$4", berjalan, is_er), _lbl(ind["status"]))
        c = ws.cell(row=n, column=6, value=ind["lalu"])
        c.number_format = fmt
        ubah = inti.selisih(ind["realisasi"], ind["lalu"])
        bk.rumus(ws, f"G{n}", f'=IF(OR(F{n}="",N(F{n})=0,B{n}=""),"",B{n}/F{n}-1)', ubah if ubah is not None else "", "+0.0%;-0.0%;0.0%")
        for col in range(1, 8):
            ws.cell(row=n, column=col).border = GARIS
        if ind["turunan"]:
            for col in range(1, 8):
                ws.cell(row=n, column=col).font = Font(bold=True)
    akhir = 6 + len(r["indikator"])
    bk.warnai_status(ws, f"E7:E{akhir}")
    ws.cell(row=akhir + 1, column=1, value="Sel kuning = target, bisa diisi/diubah. Realisasi dihitung dari lembar platform (sampai baris 3000).").font = Font(italic=True, color="7F7F7F")
    ws.freeze_panes = "A7"

    # ---------------- Per kelompok & per akun
    def lembar_kunci(nama, kol_data, baris, kunci_target):
        w = bk.lembar(nama)
        judul = [kunci_target, "Instagram", "TikTok", "Facebook", "YouTube", "Berita", "Total", "Views", "Kuota", "Capaian", "Status"]
        bk.kepala(w, 1, judul, [28, 10, 10, 10, 10, 9, 9, 13, 9, 10, 15])
        for i, b in enumerate(baris, start=2):
            w.cell(row=i, column=1, value=b["nama"])
            for j, pl in enumerate(inti.PLATFORM, start=2):
                bk.rumus(w, f"{get_column_letter(j)}{i}", f"=COUNTIF({rng[pl]}${kol_data}$2:${kol_data}${R},$A{i})", b[pl], F_ANGKA)
            bk.rumus(w, f"G{i}", f"=SUM(B{i}:F{i})", b["total"], F_ANGKA)
            bk.rumus(w, f"H{i}", "=" + "+".join(f"SUMIF({rng[pl]}${kol_data}$2:${kol_data}${R},$A{i},{rng[pl]}$H$2:$H${R})"
                                               for pl in inti.SOSMED), b["views"], F_ANGKA)
            c = w.cell(row=i, column=9, value=b["kuota"])
            c.fill = KUNING
            bk.rumus(w, f"J{i}", f'=IF(N(I{i})=0,"",G{i}/I{i})', b["capaian"] if b["capaian"] is not None else "", F_CAPAIAN)
            bk.rumus(w, f"K{i}", "=" + _status_rumus(f"J{i}", "Ringkasan!$B$4", berjalan, False), _lbl(b["status"]))
        t = len(baris) + 2
        w.cell(row=t, column=1, value="Total").font = Font(bold=True)
        for j in range(2, 10):
            col = get_column_letter(j)
            nilai = {2: "ig", 3: "tt", 4: "fb", 5: "yt", 6: "web"}
            if j <= 6:
                v = sum(b[nilai[j]] for b in baris)
            elif j == 7:
                v = sum(b["total"] for b in baris)
            elif j == 8:
                v = sum(b["views"] for b in baris)
            else:
                v = sum(b["kuota"] or 0 for b in baris)
            c = bk.rumus(w, f"{col}{t}", f"=SUM({col}2:{col}{t - 1})", v, F_ANGKA)
            c.font = Font(bold=True)
        for col in range(1, 12):
            w.cell(row=t, column=col).fill = ABU
        bk.warnai_status(w, f"K2:K{t}")
        w.freeze_panes = "B2"
        w.auto_filter.ref = f"A1:K{max(2, t - 1)}"

    if pakai_kel:
        lembar_kunci(f"Per {label}", "B", r["kelompok"], label)
    lembar_kunci("Per Akun", "D", r["akun"], "Nama Akun")

    # ---------------- lembar platform
    for pl in inti.PLATFORM:
        tulis_lembar_platform(bk, pl, _urut([p for p in posts_t if p["platform"] == pl]), label)

    # ---------------- Catatan Data
    ws = bk.lembar("Catatan Data")
    bk.kepala(ws, 1, ["Platform", "Tanggal", "Nama Akun", "Link", "Masalah"], [12, 14, 22, 50, 60])
    masalah = [p for p in _urut(posts) if flags.get(p["id"])]
    if not masalah:
        ws["A2"] = "Tidak ada masalah data pada periode ini."
    for i, p in enumerate(masalah, start=2):
        ws.cell(row=i, column=1, value=inti.NAMA_PLATFORM[p["platform"]])
        ws.cell(row=i, column=2, value=inti.tgl_panjang(p.get("tanggal")) or "(kosong)")
        ws.cell(row=i, column=3, value=p.get("akun") or "(kosong)")
        ws.cell(row=i, column=4, value=p.get("link") or "(kosong)")
        ws.cell(row=i, column=5, value=", ".join(inti.FLAG[f][0] for f in flags[p["id"]]))
    ws.freeze_panes = "A2"
    return bk.simpan()


def buat_data(klien, ym, rows, flags, filter_teks):
    """Unduh data terfilter: satu lembar Data, baris total SUBTOTAL(109, ...) yang mengikuti filter Excel."""
    bk = Buku(klien.get("warna"))
    label = klien.get("label_kelompok") or "Kelompok"
    ws = bk.lembar("Data")
    ws["A1"] = f"{klien['nama']} — {inti.nama_periode(ym)}"
    ws["A1"].font = Font(bold=True, size=14, color=bk.warna)
    ws["A2"] = "Filter: " + filter_teks
    ws["A2"].font = Font(italic=True)
    judul = ["No", "Platform", label, "Tanggal", "Nama Akun", "Jenis", "Judul", "Link", "Views", "Like", "Comment",
             "Repost", "Share", "Interaksi", "ER", "Catatan"]
    bk.kepala(ws, 4, judul, [5, 11, 18, 12, 20, 9, 46, 38, 11, 9, 9, 9, 9, 11, 9, 36])
    awal = 5
    for i, p in enumerate(rows, start=awal):
        ws.cell(row=i, column=1, value=i - awal + 1)
        ws.cell(row=i, column=2, value=inti.NAMA_PLATFORM[p["platform"]])
        ws.cell(row=i, column=3, value=p.get("kelompok") or None)
        if p.get("tanggal"):
            ws.cell(row=i, column=4, value=date.fromisoformat(p["tanggal"])).number_format = "dd/mm/yyyy"
        ws.cell(row=i, column=5, value=p.get("akun") or None)
        ws.cell(row=i, column=6, value=p.get("jenis"))
        ws.cell(row=i, column=7, value=p.get("judul") or None)
        c = ws.cell(row=i, column=8, value=p.get("link") or None)
        info = inti.parse_link(p.get("link"))
        if info:
            c.hyperlink = info["url"]
            c.font = Font(color="0563C1", underline="single")
        if p["platform"] != "web":
            for j, m in enumerate(inti.METRIK, start=9):
                ws.cell(row=i, column=j, value=p.get(m)).number_format = F_ANGKA
            bk.rumus(ws, f"N{i}", f"=SUM(J{i}:M{i})", inti.interaksi(p), F_ANGKA)
            er = inti.er_baris(p)
            bk.rumus(ws, f"O{i}", f'=IFERROR(N{i}/I{i},"")', er if er is not None else "", F_PERSEN)
        ws.cell(row=i, column=16, value=", ".join(inti.FLAG[f][0] for f in flags.get(p["id"], [])) or None)
    akhir = awal + len(rows) - 1
    t = akhir + 1
    ws.cell(row=t, column=1, value="Total").font = Font(bold=True)
    tot = {m: sum(int(p.get(m) or 0) for p in rows if p["platform"] != "web") for m in inti.METRIK}
    it = sum(inti.interaksi(p) for p in rows)
    for j, m in enumerate(inti.METRIK + ["interaksi"], start=9):
        col = get_column_letter(j)
        bk.rumus(ws, f"{col}{t}", f"=SUBTOTAL(109,{col}{awal}:{col}{max(awal, akhir)})", tot.get(m, it), F_ANGKA).font = Font(bold=True)
    bk.rumus(ws, f"O{t}", f'=IFERROR(N{t}/I{t},"")', (it / tot["views"]) if tot["views"] else "", F_PERSEN).font = Font(bold=True)
    for col in range(1, 17):
        ws.cell(row=t, column=col).fill = ABU
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:P{max(awal, akhir)}"
    return bk.simpan()
