"""
Impor workbook lama (.xlsx) ke satu periode klien.

Lembar dikenali dari namanya (IG/Instagram, TikTok, Facebook, YouTube, Artikel/Berita). Kolom dibaca dari judulnya.
Hanya baris bertanggal di periode aktif yang masuk; tanggal salah tahun pada bulan yang sama dikoreksi dan dilaporkan.

    python scripts/impor_workbook.py --cek  file.xlsx 2026-08                 # hanya periksa, tidak menyimpan
    python scripts/impor_workbook.py --klien <id> file.xlsx 2026-08           # simpan ke klien
    python scripts/impor_workbook.py --template contoh_template.xlsx          # buat template kosong
"""
import argparse
import json
import os
import sys
import warnings

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import inti  # noqa: E402

LEMBAR = [("ig", ("ig", "instagram")), ("tt", ("tiktok", "tt")), ("fb", ("facebook", "fb")),
          ("yt", ("youtube", "yt")), ("web", ("artikel", "berita", "web", "portal", "media online"))]


def kenali_lembar(nama):
    k = inti.kunci_kolom(nama)
    if k.startswith("pivot"):
        return None
    for pl, nama_dikenal in LEMBAR:
        for n in nama_dikenal:
            nk = inti.kunci_kolom(n)
            if k == nk or (len(nk) > 2 and k.startswith(nk)):
                return pl
    return None


def baca_workbook(sumber, ym, alias=None, hari_ini=None):
    """-> (rows, laporan). rows siap untuk inti.gabung(pertahankan_ganda=True)."""
    from openpyxl import load_workbook
    alias = alias or {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(sumber, data_only=True, read_only=True)
    rows, laporan = [], {"lembar": [], "diabaikan": [], "dikoreksi": []}
    for ws in wb.worksheets:
        pl = kenali_lembar(ws.title)
        if not pl:
            laporan["diabaikan"].append(ws.title)
            continue
        semua = list(ws.iter_rows(values_only=True))
        kepala_i = next((i for i, r in enumerate(semua[:15]) if "link" in inti.peta_kolom(r)), None)
        info = {"nama": ws.title, "platform": pl, "masuk": 0, "luar_periode": 0, "tanpa_tanggal": 0,
                "tanpa_link": 0, "dikoreksi": 0}
        if kepala_i is None:
            info["galat"] = "kolom Link tidak ditemukan"
            laporan["lembar"].append(info)
            continue
        peta = inti.peta_kolom(semua[kepala_i])
        for n, r in enumerate(semua[kepala_i + 1:], start=kepala_i + 2):
            ambil = lambda f: r[peta[f]] if f in peta and peta[f] < len(r) else None  # noqa: E731
            link = str(ambil("link") or "").strip()
            if not link:
                if any(ambil(f) not in (None, "") for f in ("judul", "akun", "tanggal")):
                    info["tanpa_link"] += 1
                continue
            tgl = inti.baca_tanggal(ambil("tanggal"))
            if not tgl:
                info["tanpa_tanggal"] += 1
                continue
            tgl2, koreksi = inti.koreksi_tahun(tgl, ym)
            if tgl2[:7] != ym:
                info["luar_periode"] += 1
                continue
            if koreksi:
                info["dikoreksi"] += 1
                laporan["dikoreksi"].append(f"{ws.title} baris {n}: {inti.tgl_panjang(tgl)} → {inti.tgl_panjang(tgl2)}")
            row = {"platform": pl, "link": link, "tanggal": tgl2,
                   "kelompok": inti.terapkan_alias(str(ambil("kelompok") or "").strip(), alias.get("kelompok", {})),
                   "akun": inti.terapkan_alias(str(ambil("akun") or "").strip(), alias.get("akun", {})),
                   "jenis": inti.jenis_dari(ambil("jenis"), pl), "judul": str(ambil("judul") or "").strip()}
            for m in inti.METRIK:
                row[m] = inti.baca_angka(ambil(m))
            rows.append(row)
            info["masuk"] += 1
        laporan["lembar"].append(info)
    wb.close()
    laporan["total"] = len(rows)
    return rows, laporan


def buat_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    kolom = ["No", "DPC", "Tanggal", "Nama Akun", "Video/Flyer", "Judul", "Link", "Views", "Like", "Comment", "Repost", "Share"]
    for nama in ("IG", "TikTok", "Facebook", "YouTube", "Berita"):
        ws = wb.create_sheet(nama)
        ws.append(kolom if nama != "Berita" else ["No", "DPC", "Tanggal", "Nama Akun", "Jenis", "Judul", "Link"])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFE699")
        ws.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", nargs="?")
    ap.add_argument("periode", nargs="?", help="YYYY-MM")
    ap.add_argument("--cek", action="store_true", help="hanya periksa, tidak menyimpan")
    ap.add_argument("--klien", help="id klien tujuan")
    ap.add_argument("--template", metavar="XLSX", help="buat template workbook kosong")
    a = ap.parse_args()
    if a.template:
        buat_template(a.template)
        print("Template dibuat:", a.template)
        return
    if not a.file or not a.periode:
        ap.error("file dan periode wajib diisi")
    import penyimpanan as ps
    alias = {}
    if a.klien:
        k = ps.ambil_klien(a.klien)
        if not k:
            sys.exit(f"Klien {a.klien} tidak ditemukan. Klien yang ada: " + ", ".join(f"{x['id']} ({x['nama']})" for x in ps.semua_klien()))
        alias = k.get("alias") or {}
    rows, lap = baca_workbook(a.file, a.periode, alias)
    for l in lap["lembar"]:
        print(f"- {l['nama']} ({inti.NAMA_PLATFORM[l['platform']]}): {l['masuk']} masuk, {l['luar_periode']} di luar periode, "
              f"{l['tanpa_tanggal']} tanpa tanggal, {l['dikoreksi']} tanggal dikoreksi{'  ' + l['galat'] if l.get('galat') else ''}")
    if lap["diabaikan"]:
        print("Lembar tidak dikenali (dilewati):", ", ".join(lap["diabaikan"]))
    for d in lap["dikoreksi"]:
        print("  koreksi:", d)
    print("Total baris masuk:", lap["total"])
    if a.cek or not a.klien:
        print("(mode cek: tidak ada yang disimpan)")
        return
    with ps.KUNCI:
        per = ps.ambil_periode(a.klien, a.periode, buat=True)
        baru, upd = inti.gabung(per["posts"], rows, lambda: ps.id_baru("p"), pertahankan_ganda=True)
        ps.simpan_periode(a.klien, per)
    print(f"Disimpan: {baru} baru, {upd} diperbarui.")


if __name__ == "__main__":
    main()
