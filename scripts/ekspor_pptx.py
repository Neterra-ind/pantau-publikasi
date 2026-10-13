"""Ekspor PPTX laporan bulanan (16:9). Warna dan logo mengikuti klien; grafik berupa grafik asli PowerPoint
(data bisa diedit); link konten berupa hyperlink. Semua angka berasal dari inti.ringkas (sama dengan layar)."""
import calendar
import os
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import inti

FONT = "Arial"
GELAP = "1A1A1A"
REDUP = "6B6B6B"
GARIS = "E3E3E3"
WARNA_STATUS = {"tercapai": ("C6EFCE", "0F5323"), "sesuai": ("DDEBF7", "1F4E79"), "hampir": ("FFEB9C", "7F6000"),
                "dikejar": ("FCE4D6", "843C0C"), "tertinggal": ("FFC7CE", "9C0006"), "belum": ("EDEDED", "404040")}
WARNA_PL = {"ig": "C13584", "tt": "111111", "fb": "1877F2", "yt": "E00000", "web": "6B7A8F"}
WARNA_JENIS = {"Video": None, "Flyer": "F2A900", "Artikel": "6B7A8F"}


def rgb(h):
    return RGBColor.from_string(h.upper())


def gelapkan(h, f=0.55):
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return "%02X%02X%02X" % (int(r * f), int(g * f), int(b * f))


def potong(t, n):
    t = str(t or "")
    return t if len(t) <= n else t[: n - 1].rstrip() + "…"


class Deck:
    def __init__(self, klien, r, logo=None):
        self.k, self.r = klien, r
        self.warna = (klien.get("warna") or "C00000").lstrip("#").upper()
        self.logo = logo if logo and os.path.exists(logo) else None
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        y, m = int(r["ym"][:4]), int(r["ym"][5:7])
        self.rentang = f"1–{calendar.monthrange(y, m)[1]} {inti.BULAN[m - 1]} {y}"
        self.label = klien.get("label_kelompok") or "Kelompok"
        self.nomor = 0

    # ------------------------------------------------------------ dasar
    def baru(self, judul=None, sub=None, latar="FFFFFF"):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(latar)
        self.nomor += 1
        if judul:
            self.kotak(s, 0, 0, 0.18, 7.5, self.warna)
            self.teks(s, 0.6, 0.35, 10.6, 0.7, judul, 28, True)
            self.teks(s, 0.6, 0.98, 10.6, 0.4, sub or f"{self.k['nama']} · {self.r['nama_periode']}", 13, False, REDUP)
            self.logo_di(s, 11.75, 0.3, 1.2, 0.9, gelap=False)
            self.teks(s, 12.3, 7.05, 0.8, 0.3, str(self.nomor), 10, False, REDUP, PP_ALIGN.RIGHT)
        return s

    def kotak(self, s, x, y, w, h, isi="FFFFFF", garis=None, bentuk=MSO_SHAPE.RECTANGLE):
        k = s.shapes.add_shape(bentuk, Inches(x), Inches(y), Inches(w), Inches(h))
        k.shadow.inherit = False
        if isi:
            k.fill.solid()
            k.fill.fore_color.rgb = rgb(isi)
        else:
            k.fill.background()
        if garis:
            k.line.color.rgb = rgb(garis)
            k.line.width = Pt(1)
        else:
            k.line.fill.background()
        if bentuk == MSO_SHAPE.ROUNDED_RECTANGLE:
            k.adjustments[0] = 0.08
        return k

    def teks(self, s, x, y, w, h, baris, ukuran=14, tebal=False, warna=GELAP, rata=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, url=None, spasi=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, b in enumerate([baris] if isinstance(baris, str) else baris):
            t, u, tb_, w_ = (b, ukuran, tebal, warna) if isinstance(b, str) else (tuple(b) + (ukuran, tebal, warna)[len(b) - 1:])[:4]
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = rata
            if spasi:
                p.space_after = Pt(spasi)
            run = p.add_run()
            run.text = t
            run.font.name, run.font.size, run.font.bold = FONT, Pt(u), tb_
            run.font.color.rgb = rgb(w_)
            if url:
                run.hyperlink.address = url
        return tb

    def logo_di(self, s, x, y, w, h, gelap=True):
        if self.logo:
            from PIL import Image
            iw, ih = Image.open(self.logo).size
            k = min(w / iw, h / ih)
            pw, ph = iw * k, ih * k
            s.shapes.add_picture(self.logo, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        elif gelap:
            self.teks(s, x, y, w, h, self.k["nama"].upper(), 14, True, "FFFFFF", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    def kartu(self, s, x, y, w, h, label, nilai, bawah=None, warna_bawah=REDUP):
        self.kotak(s, x, y, w, h, "FFFFFF", GARIS, MSO_SHAPE.ROUNDED_RECTANGLE)
        self.kotak(s, x, y + 0.12, 0.07, h - 0.24, self.warna)
        self.teks(s, x + 0.25, y + 0.14, w - 0.4, 0.3, label.upper(), 10.5, True, REDUP)
        self.teks(s, x + 0.25, y + 0.42, w - 0.4, 0.55, nilai, 24 if len(nilai) < 12 else 20, True)
        if bawah:
            self.teks(s, x + 0.25, y + h - 0.42, w - 0.4, 0.3, bawah, 10.5, False, warna_bawah)

    def tabel(self, s, x, y, lebar, kepala, baris, ukuran=11, tinggi=0.32, rata=None, total=None, status_kol=None,
              link_kol=None):
        nr, nc = len(baris) + 1 + (1 if total else 0), len(kepala)
        gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(lebar)), Inches(tinggi * nr))
        t = gs.table
        t.first_row, t.horz_banding = True, False
        for i, w in enumerate(lebar):
            t.columns[i].width = Inches(w)
        rata = rata or [PP_ALIGN.LEFT] * nc

        def isi(c, teks_, fill, warna, tebal=False, rt=PP_ALIGN.LEFT, url=None, u=ukuran):
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(fill)
            c.margin_left = c.margin_right = Inches(0.07)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            c.text_frame.word_wrap = True
            p.alignment = rt
            run = p.add_run()
            run.text = str(teks_)
            run.font.name, run.font.size, run.font.bold = FONT, Pt(u), tebal
            run.font.color.rgb = rgb(warna)
            if url:
                run.hyperlink.address = url

        for j, h in enumerate(kepala):
            isi(t.cell(0, j), h, self.warna, "FFFFFF", True, rata[j])
        for i, b in enumerate(baris, start=1):
            fill = "FFFFFF" if i % 2 else "F5F5F5"
            for j, v in enumerate(b):
                if status_kol is not None and j == status_kol and isinstance(v, tuple):
                    st = v[1]
                    bg, fg = WARNA_STATUS.get(st, (fill, GELAP))
                    isi(t.cell(i, j), v[0], bg, fg, True, rata[j])
                elif link_kol is not None and j == link_kol and isinstance(v, tuple):
                    isi(t.cell(i, j), v[0], fill, "0B57D0", False, rata[j], v[1], ukuran - 1)
                else:
                    isi(t.cell(i, j), v, fill, GELAP, False, rata[j])
        if total:
            for j, v in enumerate(total):
                isi(t.cell(nr - 1, j), v, gelapkan(self.warna, 0.75), "FFFFFF", True, rata[j])
        for row in t.rows:
            row.height = Inches(tinggi)
        return gs

    def _judul_chart(self, ch, judul):
        ch.font.size, ch.font.name = Pt(9), FONT
        if judul:
            ch.has_title = True
            ch.chart_title.text_frame.text = judul
            run = ch.chart_title.text_frame.paragraphs[0].runs[0]
            run.font.size, run.font.bold = Pt(12), True
        else:
            ch.has_title = False

    def batang(self, s, x, y, w, h, kategori, seri, judul=None, horizontal=False, tumpuk=False, warna=None, format_angka="#,##0"):
        if not kategori:
            self.teks(s, x, y + h / 2 - 0.2, w, 0.4, f"{judul or 'Grafik'}: belum ada data (nama akun kosong).", 11, False, REDUP, PP_ALIGN.CENTER)
            return None
        cd = CategoryChartData()
        cd.categories = kategori
        for nama, nilai in seri:
            cd.add_series(nama, nilai)
        tipe = (XL_CHART_TYPE.BAR_STACKED if tumpuk else XL_CHART_TYPE.BAR_CLUSTERED) if horizontal else \
            (XL_CHART_TYPE.COLUMN_STACKED if tumpuk else XL_CHART_TYPE.COLUMN_CLUSTERED)
        ch = s.shapes.add_chart(tipe, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        self._judul_chart(ch, judul)
        ch.has_legend = len(seri) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(9)
        pl = ch.plots[0]
        pl.gap_width = 55
        if tumpuk:
            pl.overlap = 100
        pl.has_data_labels = not tumpuk
        if not tumpuk:
            pl.data_labels.font.size = Pt(8)
            pl.data_labels.number_format = format_angka
            pl.data_labels.number_format_is_linked = False
            pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        for i, ser in enumerate(pl.series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb((warna or [self.warna, "F2A900", "6B7A8F"])[i % 3])
        ch.value_axis.major_gridlines.format.line.color.rgb = rgb("E6E6E6")
        ch.value_axis.format.line.fill.background()
        ch.value_axis.tick_labels.font.size = Pt(8)
        ch.value_axis.tick_labels.number_format = format_angka
        ch.value_axis.tick_labels.number_format_is_linked = False
        ch.category_axis.tick_labels.font.size = Pt(8)
        if horizontal:
            ch.category_axis.reverse_order = True
        return ch

    def cincin(self, s, x, y, w, h, pasangan, judul=None, warna=None):
        cd = CategoryChartData()
        cd.categories = [k for k, _ in pasangan]
        cd.add_series("Jumlah", [v for _, v in pasangan])
        ch = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        self._judul_chart(ch, judul)
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
        pl = ch.plots[0]
        pl.has_data_labels = True
        pl.data_labels.show_percentage, pl.data_labels.show_value = True, False
        pl.data_labels.number_format = "0%"
        pl.data_labels.number_format_is_linked = False
        pl.data_labels.font.size = Pt(10)
        pl.data_labels.font.bold = True
        pl.data_labels.font.color.rgb = rgb("FFFFFF")
        for i in range(len(pasangan)):
            pt = pl.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb((warna or [self.warna, "F2A900", "6B7A8F"])[i % 3])
        return ch

    # ------------------------------------------------------------ slide
    def sampul(self, penutup=False):
        s = self.baru(latar="111111")
        self.nomor -= 1
        o = self.kotak(s, 7.9, -1.2, 7.8, 7.8, self.warna, bentuk=MSO_SHAPE.OVAL)
        o.fill.gradient()
        o.fill.gradient_angle = 315
        o.fill.gradient_stops[0].color.rgb = rgb(gelapkan(self.warna, 0.4))
        o.fill.gradient_stops[1].color.rgb = rgb(self.warna)
        self.kotak(s, 7.4, 5.4, 1.1, 1.1, gelapkan(self.warna, 0.6), bentuk=MSO_SHAPE.OVAL)
        self.logo_di(s, 9.6, 1.6, 3.0, 2.8)
        if penutup:
            self.teks(s, 0.9, 2.4, 7.0, 1.2, "TERIMA KASIH", 52, True, "FFFFFF")
            self.teks(s, 0.9, 3.8, 6.6, 1.0, [(self.k["nama"], 18, True, "FFFFFF"), (f"Periode {self.rentang}", 15, False, "BDBDBD")])
        else:
            self.teks(s, 0.9, 1.7, 6.8, 2.2, ["LAPORAN", "PUBLIKASI KONTEN"], 44, True, "FFFFFF")
            self.teks(s, 0.9, 4.25, 6.6, 0.5, "PERIODE " + self.rentang.upper(), 18, False, "FFFFFF")
            self.teks(s, 0.9, 4.85, 6.6, 0.9, self.k["nama"], 15, False, "BDBDBD")

    def ringkasan(self):
        r = self.r
        s = self.baru("Ringkasan")
        t, lalu = r["total"], (r["lalu"] or {}).get("total")

        def beda(k):
            if not lalu:
                return None, REDUP
            d = inti.selisih(t[k], lalu[k])
            if d is None:
                return None, REDUP
            return f"{'▲' if d >= 0 else '▼'} {inti.persen(abs(d), 1)} vs bulan lalu", "1E7B34" if d >= 0 else "B3261E"
        for i, (k, label, nilai) in enumerate([("publikasi", "Publikasi", inti.ribu(t["publikasi"])),
                                               ("views", "Views", inti.ribu(t["views"])),
                                               ("interaksi", "Interaksi", inti.ribu(t["interaksi"])),
                                               ("er", "Engagement rate", inti.persen(t["er"]))]):
            b, w = beda(k)
            self.kartu(s, 0.6 + i * 3.08, 1.6, 2.88, 1.35, label, nilai, b, w)
        kepala = ["Platform", "Postingan", "Video", "Flyer", "Artikel", "Views", "Interaksi", "ER"]
        baris = []
        for pl in inti.PLATFORM:
            x = r["per"][pl]
            if pl == "web":
                baris.append([inti.NAMA_PLATFORM[pl], inti.ribu(x["posts"]), inti.ribu(x["video"]), inti.ribu(x["flyer"]),
                              inti.ribu(x["artikel"]), "–", "–", "–"])
            else:
                baris.append([inti.NAMA_PLATFORM[pl], inti.ribu(x["posts"]), inti.ribu(x["video"]), inti.ribu(x["flyer"]),
                              inti.ribu(x["artikel"]), inti.ribu(x["views"]), inti.ribu(x["interaksi"]), inti.persen(x["er"])])
        total = ["Total", inti.ribu(t["publikasi"]), inti.ribu(t["video"]), inti.ribu(t["flyer"]), inti.ribu(t["artikel"]),
                 inti.ribu(t["views"]), inti.ribu(t["interaksi"]), inti.persen(t["er"])]
        R = PP_ALIGN.RIGHT
        self.tabel(s, 0.6, 3.35, [2.6, 1.4, 1.2, 1.2, 1.2, 1.7, 1.5, 1.3], kepala, baris, 12, 0.42,
                   [PP_ALIGN.LEFT] + [R] * 7, total)
        self.teks(s, 0.6, 6.55, 12, 0.4, "Views, interaksi, dan ER hanya media sosial; berita tidak memiliki views. "
                  "ER = (like + comment + repost + share) ÷ views.", 10, False, REDUP)

    def target(self):
        ind = [i for i in self.r["indikator"] if i["target"]]
        if not ind:
            return
        sk = self.r["skor"]
        for n in range(0, len(ind), 12):
            s = self.baru("Pencapaian target", f"{sk['tercapai']} dari {sk['dinilai']} target tercapai · "
                          + ("periode berjalan, status memakai laju harian" if self.r["berjalan"] else "periode sudah lewat"))
            baris = []
            for i in ind[n:n + 12]:
                f = inti.persen if i["jenis"] == "er" else inti.ribu
                baris.append([i["label"], f(i["realisasi"]), f(i["target"]),
                              inti.persen(i["capaian"], 0) if i["capaian"] is not None else "–",
                              (inti.STATUS_LABEL[i["status"]], i["status"]) if i["status"] else "–"])
            R = PP_ALIGN.RIGHT
            self.tabel(s, 0.6, 1.6, [4.4, 2.0, 2.0, 1.6, 2.1], ["Indikator", "Realisasi", "Target", "Capaian", "Status"],
                       baris, 12, 0.4, [PP_ALIGN.LEFT, R, R, R, PP_ALIGN.CENTER], status_kol=4)

    def kelompok(self):
        kel = [b for b in self.r["kelompok"] if b["total"] or b["kuota"]]
        if not kel:
            return
        s = self.baru(f"Sebaran per {self.label}", f"{len([b for b in kel if b['total']])} {self.label} dengan publikasi"
                      + (f" · kuota terpenuhi {self.r['kuota_kelompok']['terpenuhi']} dari {self.r['kuota_kelompok']['dinilai']}"
                         if self.r["kuota_kelompok"]["dinilai"] else ""))
        top = sorted(kel, key=lambda b: (-b["total"], b["nama"].lower()))[:15]
        seri = [(inti.NAMA_PLATFORM[pl], [b[pl] for b in top]) for pl in inti.PLATFORM if any(b[pl] for b in top)]
        self.batang(s, 0.5, 1.5, 12.3, 5.6, [potong(b["nama"], 22) for b in top], seri,
                    f"15 {self.label} dengan publikasi terbanyak", tumpuk=True, warna=[WARNA_PL[pl] for pl in inti.PLATFORM
                                                                                        if any(b[pl] for b in top)])
        urut = sorted(kel, key=lambda b: b["nama"].lower())
        per_kolom = 13
        ada_kuota = any(b["kuota"] for b in urut)
        for n in range(0, len(urut), per_kolom * 2):
            s = self.baru(f"Rincian per {self.label}", "Urut abjad" + (" · kuota dan status" if ada_kuota else ""))
            for c in range(2):
                bagian = urut[n + c * per_kolom: n + (c + 1) * per_kolom]
                if not bagian:
                    continue
                if ada_kuota:
                    baris = [[potong(b["nama"], 26), inti.ribu(b["total"]), inti.ribu(b["kuota"]) if b["kuota"] else "–",
                              (inti.STATUS_LABEL[b["status"]], b["status"]) if b["status"] else "–"] for b in bagian]
                    self.tabel(s, 0.5 + c * 6.25, 1.5, [2.75, 1.0, 0.9, 1.4], [self.label, "Posting", "Kuota", "Status"],
                               baris, 10.5, 0.38, [PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER], status_kol=3)
                else:
                    baris = [[potong(b["nama"], 26), inti.ribu(b["total"]), inti.ribu(b["views"])] for b in bagian]
                    self.tabel(s, 0.5 + c * 6.25, 1.5, [3.25, 1.3, 1.5], [self.label, "Posting", "Views"], baris, 10.5, 0.38,
                               [PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT])

    def platform(self, pl):
        r = self.r
        x = r["per"][pl]
        d = r["detail_platform"].get(pl)
        if not d or not x["views"]:
            return
        ind = {i["kunci"]: i for i in r["indikator"]}
        s = self.baru(f"Detail {inti.NAMA_PLATFORM[pl]}")
        ip, ie = ind.get("posts." + pl), ind.get("er." + pl)
        self.kartu(s, 0.6, 1.5, 2.88, 1.2, "Postingan", inti.ribu(x["posts"]),
                   f"Target {inti.ribu(ip['target'])} · {inti.persen(ip['capaian'], 0)}" if ip and ip["target"] else None)
        self.kartu(s, 3.68, 1.5, 2.88, 1.2, "Views", inti.ribu(x["views"]))
        self.kartu(s, 6.76, 1.5, 2.88, 1.2, "Interaksi", inti.ribu(x["interaksi"]))
        self.kartu(s, 9.84, 1.5, 2.88, 1.2, "ER", inti.persen(x["er"]),
                   f"Target min. {inti.persen(ie['target'])}" if ie and ie["target"] else None)
        aktif = d["aktif"][:7]
        views = d["views"][:7]
        self.batang(s, 0.5, 2.9, 6.1, 2.55, [potong(a["nama"], 20) for a in aktif], [("Postingan", [a["posts"] for a in aktif])],
                    "Akun paling aktif", horizontal=True)
        self.batang(s, 6.75, 2.9, 6.1, 2.55, [potong(a["nama"], 20) for a in views], [("Views", [a["views"] for a in views])],
                    "Akun dengan views tertinggi", horizontal=True, warna=[gelapkan(self.warna, 0.6)])
        self.teks(s, 0.6, 5.55, 6, 0.3, "KONTEN VIEWS TERTINGGI", 10.5, True, REDUP)
        for i, k in enumerate(d["konten"]):
            y = 5.85 + i * 0.4
            self.teks(s, 0.6, y, 0.35, 0.35, str(i + 1), 12, True, self.warna)
            self.teks(s, 0.95, y, 8.6, 0.35, potong(k["judul"] or "(tanpa judul)", 95), 11, False, "0B57D0",
                      url=(inti.parse_link(k["link"]) or {}).get("url"))
            self.teks(s, 9.6, y, 3.2, 0.35, f"{potong(k['akun'], 18)} · {inti.ribu(k['views'])} views", 11, True, GELAP, PP_ALIGN.RIGHT)

    def berita(self, posts):
        rows = sorted([p for p in posts if p["platform"] == "web"], key=lambda p: (p.get("tanggal") or "9", p.get("akun") or ""))
        for n in range(0, len(rows), 8):
            s = self.baru("Berita", f"{len(rows)} artikel berita · {self.r['nama_periode']}")
            baris = []
            for i, p in enumerate(rows[n:n + 8], start=n + 1):
                url = (inti.parse_link(p.get("link")) or {}).get("url")
                baris.append([str(i), inti.tgl_panjang(p.get("tanggal")) or "–", potong(p.get("akun"), 22),
                              potong(p.get("judul") or "(tanpa judul)", 80), (potong(p.get("link"), 42), url) if url else potong(p.get("link"), 42)])
            self.tabel(s, 0.5, 1.5, [0.5, 1.8, 2.2, 5.1, 2.7], ["No", "Tanggal", "Sumber", "Judul", "Link"], baris, 10.5, 0.6,
                       link_kol=4)

    def produksi(self):
        r, t = self.r, self.r["total"]
        s = self.baru("Produksi konten")
        for i, (j, warna) in enumerate([("Video", self.warna), ("Flyer", "F2A900"), ("Artikel", "6B7A8F")]):
            self.kartu(s, 0.6, 1.55 + i * 1.45, 3.0, 1.25, j, inti.ribu(t[j.lower()]),
                       inti.persen(t[j.lower()] / t["publikasi"], 1) + " dari publikasi" if t["publikasi"] else None)
        pasangan = [(j, t[j.lower()]) for j in inti.JENIS if t[j.lower()]]
        if pasangan:
            self.cincin(s, 3.85, 1.4, 3.9, 4.9, pasangan, "Komposisi jenis",
                        [{"Video": self.warna, "Flyer": "F2A900", "Artikel": "6B7A8F"}[k] for k, _ in pasangan])
        kat = [inti.NAMA_PLATFORM[pl] for pl in inti.PLATFORM]
        self.batang(s, 7.9, 1.4, 5.0, 4.9, kat, [(j, [r["per"][pl][j.lower()] for pl in inti.PLATFORM]) for j in inti.JENIS],
                    "Jenis konten per platform", tumpuk=True)


def buat(klien, r, posts, logo=None):
    d = Deck(klien, r, logo)
    d.sampul()
    d.ringkasan()
    d.target()
    if klien.get("pakai_kelompok", True):
        d.kelompok()
    for pl in inti.SOSMED:
        d.platform(pl)
    d.berita(posts)
    d.produksi()
    d.sampul(penutup=True)
    buf = BytesIO()
    d.prs.save(buf)
    return buf.getvalue()
